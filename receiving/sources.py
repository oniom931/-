"""발주현황 / 시트3 읽기·쓰기. 구글 시트 API(서비스 계정)와 xlsx 파일(오프라인 테스트용) 둘 다 지원."""
from __future__ import annotations

import base64
import datetime as dt
import json
import os
import re
from dataclasses import dataclass, field
from urllib.parse import quote

from .core import OrderRow, to_date

ORDER_SPREADSHEET_ID = os.environ.get("ORDER_SPREADSHEET_ID", "1y5FAyvMTR8XJbUramt5RnnJWGqnzQopFjK6Gq04IlB0")
ORDER_SHEET = os.environ.get("ORDER_SHEET", "일간 일정")
SCHEDULE_SPREADSHEET_ID = os.environ.get("SCHEDULE_SPREADSHEET_ID", "129PazM4GUk_nkK8phXvlzWoHkvlGvGX7UOGkS1C5YWs")
SCHEDULE_SHEET = os.environ.get("SCHEDULE_SHEET", "시트3")

# 발주현황 열 위치 (0부터): C=모품목, I=발주일, J=입고예정일, K=입고일, L=품번, M=품명, R=부족, W=비고
COL = dict(product=2, order_date=8, expected=9, received=10, part_no=11, part_name=12, shortage=17, remark=22)
ORDER_FIRST_ROW = 4  # 1~3행은 제목/요약/헤더

CALENDAR_LABEL = "원부자재"  # A열이 이 글자로 시작하는 행이 입고 일정을 적는 행
CAL_COLS = "BCDEFGH"         # 일~토


def _row_to_order(row_no: int, v: list) -> OrderRow:
    get = lambda k: v[COL[k]] if len(v) > COL[k] else None
    return OrderRow(row_no, get("product"), get("part_no"), get("part_name"), get("expected"),
                    get("received"), get("order_date"), get("shortage"), get("remark"))


@dataclass
class Calendar:
    start: dt.date                       # B1 (첫 주 일요일)
    rows: list[int]                      # '원부자재 입출고' 행 번호들
    values: dict[str, str]               # 셀 주소 → 현재 내용
    skip: set[str] = field(default_factory=set)  # 병합 셀의 가운데/아래쪽 (쓸 수 없음)

    def cell_of(self, date: dt.date) -> str | None:
        diff = (date - self.start).days
        if diff < 0:
            return None
        week, wday = divmod(diff, 7)
        candidates = [r for r in self.rows if (r - 2) % 5 == 0 and (r - 2) // 5 == week]
        if not candidates:
            return None
        addr = f"{CAL_COLS[wday]}{candidates[0]}"
        return None if addr in self.skip else addr


def _calendar(b1, a_column: list, values: dict, merges: list[tuple[int, int, int, int]]) -> Calendar:
    start = to_date(b1)
    if start is None:
        raise ValueError(f"{SCHEDULE_SHEET}!B1 에서 시작 날짜를 읽지 못했습니다: {b1!r}")
    rows = [i + 1 for i, a in enumerate(a_column) if str(a or "").strip().startswith(CALENDAR_LABEL)]
    skip = set()
    for r1, c1, r2, c2 in merges:  # 1부터 시작, 끝 포함
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):
                if (r, c) != (r1, c1) and 2 <= c <= 8:
                    skip.add(f"{CAL_COLS[c - 2]}{r}")
    return Calendar(start, rows, values, skip)


# ---------------------------------------------------------------- xlsx (테스트용)

def read_orders_xlsx(path: str) -> list[OrderRow]:
    import openpyxl
    ws = openpyxl.load_workbook(path)[ORDER_SHEET]
    out = []
    for r in ws.iter_rows(min_row=ORDER_FIRST_ROW):
        n = r[0].row
        if ws.row_dimensions[n].hidden:  # 숨김 행 제외
            continue
        out.append(_row_to_order(n, [c.value for c in r]))
    return out


def read_calendar_xlsx(path: str) -> Calendar:
    import openpyxl
    ws = openpyxl.load_workbook(path)[SCHEDULE_SHEET]
    values = {}
    for row in ws.iter_rows(min_row=1, max_col=8):
        for c in row[1:]:
            if isinstance(c.value, str):
                values[c.coordinate] = c.value
    merges = [(m.min_row, m.min_col, m.max_row, m.max_col) for m in ws.merged_cells.ranges]
    a_col = [ws.cell(i, 1).value for i in range(1, ws.max_row + 1)]
    return _calendar(ws["B1"].value, a_col, values, merges)


# ---------------------------------------------------------------- Google Sheets API

def _credentials_info() -> dict:
    raw = os.environ.get("GOOGLE_SERVICE_ACCOUNT_JSON", "").strip()
    if not raw and os.environ.get("GOOGLE_APPLICATION_CREDENTIALS"):
        with open(os.environ["GOOGLE_APPLICATION_CREDENTIALS"], encoding="utf-8") as f:
            raw = f.read().strip()
    if not raw:
        raise SystemExit("GOOGLE_SERVICE_ACCOUNT_JSON 환경 변수가 없습니다. README의 '서비스 계정 키' 절을 참고하세요.")
    if raw.startswith(("'", '"')) and raw.endswith(raw[0]):  # 따옴표로 감싸 넣은 경우
        raw = raw[1:-1]
    if not raw.startswith("{"):  # base64 로 넣은 경우
        raw = base64.b64decode(raw).decode("utf-8")
    return json.loads(raw)


class SheetsClient:
    BASE = "https://sheets.googleapis.com/v4/spreadsheets"

    def __init__(self):
        from google.auth.transport.requests import AuthorizedSession
        from google.oauth2.service_account import Credentials
        creds = Credentials.from_service_account_info(
            _credentials_info(), scopes=["https://www.googleapis.com/auth/spreadsheets"])
        self.email = creds.service_account_email
        self.session = AuthorizedSession(creds)

    def _get(self, url, **params):
        resp = self.session.get(url, params=params)
        if resp.status_code == 403:
            raise SystemExit(f"권한이 없습니다. 시트를 {self.email} 에게 공유했는지 확인하세요.\n{resp.text[:300]}")
        resp.raise_for_status()
        return resp.json()

    def values(self, spreadsheet_id, a1, render="UNFORMATTED_VALUE"):
        url = f"{self.BASE}/{spreadsheet_id}/values/{quote(a1)}"
        data = self._get(url, valueRenderOption=render, dateTimeRenderOption="SERIAL_NUMBER")
        return data.get("values", [])

    def sheet_meta(self, spreadsheet_id, sheet, fields):
        data = self._get(f"{self.BASE}/{spreadsheet_id}", ranges=_quote_sheet(sheet), fields=fields)
        return data["sheets"][0]

    def write_cells(self, spreadsheet_id, sheet, cells: dict[str, str]):
        body = {"valueInputOption": "RAW",
                "data": [{"range": f"{_quote_sheet(sheet)}!{a}", "values": [[v]]} for a, v in cells.items()]}
        resp = self.session.post(f"{self.BASE}/{spreadsheet_id}/values:batchUpdate", json=body)
        if resp.status_code == 403:
            raise SystemExit(f"쓰기 권한이 없습니다. 생산일정을 {self.email} 에게 '편집자'로 공유하세요.")
        resp.raise_for_status()
        return resp.json()


def _quote_sheet(name: str) -> str:
    return "'" + name.replace("'", "''") + "'"


def read_orders_api(client: SheetsClient) -> list[OrderRow]:
    meta = client.sheet_meta(ORDER_SPREADSHEET_ID, ORDER_SHEET,
                             "sheets(data(rowMetadata(hiddenByUser,hiddenByFilter)))")
    row_meta = meta["data"][0].get("rowMetadata", [])
    hidden = {i + 1 for i, m in enumerate(row_meta) if m.get("hiddenByUser") or m.get("hiddenByFilter")}
    rows = client.values(ORDER_SPREADSHEET_ID, f"{_quote_sheet(ORDER_SHEET)}!A1:W")
    out = []
    for i, v in enumerate(rows, start=1):
        if i < ORDER_FIRST_ROW or i in hidden:
            continue
        out.append(_row_to_order(i, v))
    return out


def read_calendar_api(client: SheetsClient) -> Calendar:
    sheet = _quote_sheet(SCHEDULE_SHEET)
    rows = client.values(SCHEDULE_SPREADSHEET_ID, f"{sheet}!A1:H120")
    values = {}
    for r, row in enumerate(rows, start=1):
        for c, v in enumerate(row[1:8], start=0):
            if isinstance(v, str) and v:
                values[f"{CAL_COLS[c]}{r}"] = v
    meta = client.sheet_meta(SCHEDULE_SPREADSHEET_ID, SCHEDULE_SHEET, "sheets(merges)")
    merges = [(m["startRowIndex"] + 1, m["startColumnIndex"] + 1, m["endRowIndex"], m["endColumnIndex"])
              for m in meta.get("merges", [])]
    a_col = [row[0] if row else None for row in rows]
    b1 = rows[0][1] if rows and len(rows[0]) > 1 else None
    return _calendar(b1, a_col, values, merges)
